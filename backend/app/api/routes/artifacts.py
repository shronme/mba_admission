from __future__ import annotations

import io
import re
import uuid

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from app.api.deps.auth import get_candidate_id_from_bearer_token
from app.core.db import get_db_session
from app.repositories.cv_draft_repository import CVDraftRepository
from app.repositories.essay_draft_repository import EssayDraftRepository

router = APIRouter(prefix="/artifacts", tags=["artifacts"])

_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")


def _safe_filename(value: str) -> str:
    v = (value or "").strip() or "artifact"
    v = re.sub(r"[^\w\-. ]+", "", v).strip()
    return (v[:80] or "artifact").replace(" ", "_")


def _markdown_to_html(body: str) -> str:
    """Render Markdown body to an HTML fragment for WeasyPrint.

    Uses the `markdown-it-py` dependency already pinned in requirements.txt.
    Falls back to an HTML-escaped `<pre>` dump if markdown-it is unavailable
    so the download endpoint never hard-fails on an import error.
    """
    try:
        from markdown_it import MarkdownIt
    except Exception:
        escaped = (
            body.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )
        return f"<pre style='white-space: pre-wrap'>{escaped}</pre>"

    md = MarkdownIt("commonmark", {"html": False, "linkify": False, "breaks": True})
    return md.render(body or "")


def _add_runs_with_bold(paragraph, text: str) -> None:
    """Add `text` to a python-docx paragraph, honoring `**bold**` spans."""
    pos = 0
    for match in _INLINE_BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _render_markdown_to_docx(doc, body: str) -> None:
    """Translate a Markdown body into python-docx paragraphs.

    Handles headings (`#`..`######`), bullet lines (`- ` / `* `), blank-line
    paragraph breaks, and inline `**bold**`. Anything else is rendered as a
    plain paragraph. This is intentionally a lightweight translator — it
    covers the subset of Markdown that `RewriteCVSignature` is prompted to
    emit, not the full CommonMark grammar.
    """
    lines = (body or "").splitlines()
    buffer: list[str] = []

    def flush_paragraph() -> None:
        if not buffer:
            return
        text = " ".join(s.strip() for s in buffer if s.strip())
        buffer.clear()
        if not text:
            return
        para = doc.add_paragraph()
        _add_runs_with_bold(para, text)

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            flush_paragraph()
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 4)
            heading_text = heading.group(2).strip()
            heading_text = _INLINE_BOLD_RE.sub(r"\1", heading_text)
            doc.add_heading(heading_text, level=level)
            continue
        bullet = _BULLET_RE.match(line)
        if bullet:
            flush_paragraph()
            para = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(para, bullet.group(1).strip())
            continue
        buffer.append(line)

    flush_paragraph()


@router.get("/{artifact_id}/download", response_model=None)
async def download_artifact(
    artifact_id: uuid.UUID,
    format: str | None = Query(default=None, description="docx|pdf"),
    candidate_id: uuid.UUID = Depends(get_candidate_id_from_bearer_token),
    session=Depends(get_db_session),
) -> StreamingResponse:
    if format not in {"docx", "pdf"}:
        raise HTTPException(status_code=400, detail="Invalid format") from None

    cv_repo = CVDraftRepository(session)
    essay_repo = EssayDraftRepository(session)

    record = await cv_repo.get_by_id_unscoped(artifact_id)
    if record is not None and getattr(record, "candidate_id", None) != candidate_id:
        raise HTTPException(status_code=403, detail="Forbidden") from None

    if record is None:
        record = await essay_repo.get_by_id_unscoped(artifact_id)
        if record is not None and getattr(record, "candidate_id", None) != candidate_id:
            raise HTTPException(status_code=403, detail="Forbidden") from None

    if record is None:
        raise HTTPException(status_code=404, detail="Artifact not found") from None

    title = getattr(record, "title", "") or "artifact"
    body = getattr(record, "body", "") or ""
    filename_base = _safe_filename(title)

    first_line = next(
        (s for s in (body or "").splitlines() if s.strip()),
        "",
    )
    body_has_leading_h1 = first_line.startswith("# ")

    if format == "docx":
        from docx import Document

        doc = Document()
        if not body_has_leading_h1:
            doc.add_heading(title, level=1)
        _render_markdown_to_docx(doc, body)
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        headers = {
            "Content-Disposition": f'attachment; filename="{filename_base}.docx"'
        }
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers,
        )

    from weasyprint import HTML

    body_html = _markdown_to_html(body)
    title_html = ""
    if not body_has_leading_h1:
        escaped_title = (
            title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )
        title_html = f"<h1>{escaped_title}</h1>"
    html_document = (
        "<!doctype html><html><head><meta charset='utf-8'>"
        "<style>"
        "@page { size: A4; margin: 20mm 18mm; }"
        "body { font-family: 'Liberation Serif', 'Noto Serif', serif;"
        " font-size: 11pt; line-height: 1.35; color: #111; }"
        "h1 { font-size: 20pt; margin: 0 0 6pt 0; }"
        "h2 { font-size: 13pt; margin: 14pt 0 4pt 0;"
        " border-bottom: 1px solid #999; padding-bottom: 2pt; }"
        "h3 { font-size: 11.5pt; margin: 10pt 0 3pt 0; }"
        "p { margin: 3pt 0; }"
        "ul { margin: 3pt 0 3pt 18pt; padding: 0; }"
        "li { margin: 1pt 0; }"
        "strong { font-weight: 600; }"
        "</style></head><body>"
        f"{title_html}{body_html}"
        "</body></html>"
    )
    pdf_bytes = HTML(string=html_document).write_pdf()
    headers = {"Content-Disposition": f'attachment; filename="{filename_base}.pdf"'}
    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers=headers,
    )

